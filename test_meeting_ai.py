"""Run explicitly: python -m unittest test_meeting_ai -v."""

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from docx import Document
from fastapi.testclient import TestClient

import app
from meeting_analysis import analyze_transcript


SAMPLE = (
    "Обсудили запуск нового проекта. Алия, подготовь отчёт до пятницы. "
    "Ерлан должен отправить договор до 25.09.2026. "
    "Нужно проверить смету. Айгүл есепті ертең дайындасын. "
    "Бекзат құжаттарды жұмаға дейін жіберсін."
)


class AnalysisTests(unittest.TestCase):
    def test_russian_and_kazakh_tasks(self):
        result = analyze_transcript(SAMPLE)
        self.assertEqual(len(result["tasks"]), 5)
        self.assertEqual(result["tasks"][0]["assignee"], "Алия")
        self.assertEqual(result["tasks"][0]["deadline"], "до пятницы")
        self.assertEqual(result["tasks"][1]["assignee"], "Ерлан")
        self.assertEqual(result["tasks"][1]["deadline"], "до 25.09.2026")
        self.assertEqual(result["tasks"][2]["assignee"], "Не определён")
        self.assertEqual(result["tasks"][2]["deadline"], "Не определён")
        self.assertEqual(result["tasks"][3]["assignee"], "Айгүл")
        self.assertEqual(result["tasks"][3]["deadline"], "ертең")
        self.assertEqual(result["tasks"][4]["deadline"], "жұмаға дейін")
        self.assertTrue(all(t["status"] == "В работе" for t in result["tasks"]))
        self.assertTrue(result["summary"])

    def test_metadata_and_split_whisper_segments(self):
        tasks = analyze_transcript(
            "Нужно подготовить\nотчёт. Ответственный: Иван Петров. Срок: до конца недели."
        )["tasks"]
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]["assignee"], "Иван Петров")
        self.assertEqual(tasks[0]["deadline"], "до конца недели")

    def test_no_fabricated_tasks_and_negation(self):
        for text in ("", "Добрый день. Спасибо за встречу.",
                     "Отчёт подготовлен вчера.", "Не нужно отправить договор.",
                     "Есепті дайындау қажет емес.", "Нужно подготовить отчёт?"):
            with self.subTest(text=text):
                self.assertEqual(analyze_transcript(text)["tasks"], [])

    def test_duplicate_and_ambiguous_owner(self):
        tasks = analyze_transcript("Я должен подготовить отчёт завтра. Я должен подготовить отчёт завтра.")["tasks"]
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]["assignee"], "Не определён")

    def test_explicit_deadline_and_non_person_subject(self):
        task = analyze_transcript("Сроки нужно согласовать. Срок: пятница.")["tasks"][0]
        self.assertEqual(task["assignee"], "Не определён")
        self.assertEqual(task["deadline"], "пятница")


class EndpointTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app.app)

    def test_analysis_and_validation(self):
        response = self.client.post("/api/analyze", json={"transcript": SAMPLE})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), analyze_transcript(SAMPLE))
        self.assertEqual(response.headers["cache-control"], "no-store")
        for body in ({}, {"transcript": None}, {"transcript": "x" * 1_000_001}):
            self.assertEqual(self.client.post("/api/analyze", json=body).status_code, 422)

    def test_docx_contains_all_sections_and_full_transcript(self):
        response = self.client.post("/api/export/docx", json={"transcript": SAMPLE})
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["content-disposition"])
        self.assertIn("wordprocessingml", response.headers["content-type"])
        document = Document(BytesIO(response.content))
        paragraphs = "\n".join(p.text for p in document.paragraphs)
        for text in ("Meeting AI", "Дата формирования:", "Краткое саммари", "Поручения", "Полный транскрипт", SAMPLE):
            self.assertIn(text, paragraphs)
        self.assertEqual(len(document.tables[0].rows), 6)
        self.assertEqual(document.tables[0].rows[1].cells[0].text, "Алия")
        self.assertEqual(document.tables[0].rows[1].cells[3].text, "В работе")

    def test_docx_empty_and_xml_characters(self):
        for text in ("", "<текст>&\x00 Айгүл есепті ертең дайындасын."):
            response = self.client.post("/api/export/docx", json={"transcript": text})
            self.assertEqual(response.status_code, 200)
            Document(BytesIO(response.content))

    def test_analysis_and_export_failures_are_isolated(self):
        with patch.object(app, "analyze_transcript", side_effect=RuntimeError()):
            self.assertEqual(self.client.post("/api/analyze", json={"transcript": SAMPLE}).status_code, 500)
        with patch.object(app, "build_docx", side_effect=ImportError()):
            self.assertEqual(self.client.post("/api/export/docx", json={"transcript": SAMPLE}).status_code, 503)
        self.assertEqual(self.client.get("/").status_code, 200)

    def test_existing_transcription_formats_and_cleanup(self):
        paths = []
        def transcribe(path, **kwargs):
            paths.append(Path(path))
            def segments():
                self.assertEqual(Path(path).read_bytes(), b"audio")
                yield SimpleNamespace(text=" Сохранённый транскрипт ")
            return segments(), None
        fake = Mock()
        fake.transcribe.side_effect = transcribe
        with patch.object(app, "model", fake):
            for extension in ("mp3", "WAV", "m4a"):
                response = self.client.post("/api/transcribe", files={"file": ("sample." + extension, b"audio")})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json(), {"transcript": "Сохранённый транскрипт"})
                self.assertFalse(paths[-1].parent.exists())
            def fail(path, **kwargs):
                paths.append(Path(path))
                def segments():
                    raise ValueError("broken audio")
                    yield
                return segments(), None
            fake.transcribe.side_effect = fail
            response = self.client.post("/api/transcribe", files={"file": ("bad.wav", b"audio")})
            self.assertEqual(response.status_code, 422)
            self.assertFalse(paths[-1].parent.exists())
        for name, data in (("empty.wav", b""), ("wrong.txt", b"audio")):
            self.assertEqual(self.client.post("/api/transcribe", files={"file": (name, data)}).status_code, 400)


if __name__ == "__main__":
    unittest.main()
