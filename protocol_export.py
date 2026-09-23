"""In-memory DOCX protocol generation. No audio, disk storage or external APIs."""

from datetime import datetime
from io import BytesIO
import re


def build_docx(transcript: str, analysis: dict) -> bytes:
    # Keep this optional dependency away from application/transcription startup.
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    def clean(text):
        return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", text)

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    document.core_properties.title = "Meeting AI — Протокол совещания"
    document.core_properties.author = "Meeting AI"
    document.add_paragraph("Meeting AI", style="Title")
    document.add_paragraph("Протокол совещания", style="Subtitle")
    document.add_paragraph("Дата формирования: " + datetime.now().astimezone().strftime("%d.%m.%Y %H:%M %z"))
    document.add_heading("Краткое саммари", level=1)
    for line in analysis["summary"].splitlines():
        document.add_paragraph(clean(line))
    document.add_heading("Поручения", level=1)
    document.add_paragraph(clean(analysis["analysis_note"]))
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [1.35, 3.05, 1.45, 1.15]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, label in zip(table.rows[0].cells, ["Ответственный", "Поручение", "Срок", "Статус"]):
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EFF3")
        cell._tc.get_or_add_tcPr().append(shading)
    repeat_header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn("w:" + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for task in analysis["tasks"]:
        for cell, key in zip(table.add_row().cells, ("assignee", "description", "deadline", "status")):
            cell.text = clean(task[key])
    if not analysis["tasks"]:
        table.add_row().cells[0].merge(table.rows[-1].cells[-1]).text = "Явные поручения не найдены."
    for row in table.rows:
        if len(row._tr.tc_lst) == 1:
            row.cells[0].width = Inches(sum(widths))
            continue
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
    document.add_heading("Полный транскрипт", level=1)
    for line in transcript.splitlines() or ["В записи нет распознанной речи."]:
        document.add_paragraph(clean(line))
    output = BytesIO()
    document.save(output)
    return output.getvalue()
